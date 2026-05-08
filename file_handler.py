# -*- coding: utf-8 -*-
"""
文件处理模块
"""

import os
import shutil
import zipfile
import tempfile
from openpyxl import load_workbook
import csv

from config import TEMP_DIR, SUPPORTED_FORMATS
from watermark import create_watermark_for_excel, add_watermark_to_sheet
from utils import get_file_type, generate_output_filename, print_progress


def process_excel_file(file_path, client_name, output_dir, watermark_config, logger):
    """
    处理单个 Excel 文件
    
    Args:
        file_path: 文件路径
        client_name: 客户名称
        output_dir: 输出目录（None 或空字符串则使用原文件目录）
        watermark_config: 水印配置
        logger: 日志记录器
        
    Returns:
        str: 输出文件路径 或 None（失败时）
    """
    try:
        # 加载工作簿
        wb = load_workbook(file_path)
        
        # 收集所有水印图片路径，保存后再清理
        watermark_paths = []
        
        # 为每个工作表添加水印
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # 创建水印图片
            watermark_path = create_watermark_for_excel(
                client_name=client_name,
                sheet_width=800,
                sheet_height=600,
                config=watermark_config
            )
            watermark_paths.append(watermark_path)
            
            # 添加水印到工作表
            add_watermark_to_sheet(ws, watermark_path)
        
        # 生成输出文件名
        output_path = generate_output_filename(file_path, client_name, output_dir)
        
        # 确保输出目录存在
        output_dir_path = os.path.dirname(output_path)
        if output_dir_path:
            os.makedirs(output_dir_path, exist_ok=True)
        else:
            os.makedirs(".", exist_ok=True)
        
        # 保存文件
        wb.save(output_path)
        wb.close()
        
        # 清理所有临时水印图片（保存后再清理）
        for wm_path in watermark_paths:
            if os.path.exists(wm_path):
                os.remove(wm_path)
        
        logger.info(f"Excel 文件处理完成: {file_path} -> {output_path}")
        return output_path
        
    except Exception as e:
        logger.error(f"处理 Excel 文件失败: {file_path}, 错误: {str(e)}")
        import traceback
        traceback.print_exc()
        return None


def process_csv_file(file_path, client_name, output_dir, watermark_config, logger):
    """
    处理 CSV 文件（转换为带水印的 Excel）
    
    Args:
        file_path: 文件路径
        client_name: 客户名称
        output_dir: 输出目录（None 或空字符串则使用原文件目录）
        watermark_config: 水印配置
        logger: 日志记录器
        
    Returns:
        str: 输出文件路径 或 None（失败时）
    """
    try:
        from openpyxl import Workbook
        
        # 创建新的工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        # 读取 CSV 数据
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader, 1):
                for col_idx, value in enumerate(row, 1):
                    ws.cell(row=row_idx, column=col_idx, value=value)
        
        # 添加水印
        watermark_path = create_watermark_for_excel(
            client_name=client_name,
            config=watermark_config
        )
        add_watermark_to_sheet(ws, watermark_path)
        
        # 生成输出文件名（CSV 转 XLSX）
        dir_name = os.path.dirname(file_path)
        file_name = os.path.basename(file_path)
        name = os.path.splitext(file_name)[0]
        target_dir = output_dir if output_dir else dir_name
        if not target_dir:
            target_dir = "."
        output_path = os.path.join(target_dir, f"{name}_for_{client_name}.xlsx")
        
        output_dir_path = os.path.dirname(output_path)
        if output_dir_path:
            os.makedirs(output_dir_path, exist_ok=True)
        else:
            os.makedirs(".", exist_ok=True)
        
        wb.save(output_path)
        wb.close()
        
        # 清理临时水印图片（保存后再清理）
        if os.path.exists(watermark_path):
            os.remove(watermark_path)
        
        logger.info(f"CSV 文件处理完成: {file_path} -> {output_path}")
        return output_path
        
    except Exception as e:
        logger.error(f"处理 CSV 文件失败: {file_path}, 错误: {str(e)}")
        import traceback
        traceback.print_exc()
        return None


def process_archive(archive_path, client_name, output_dir, watermark_config, logger):
    """
    处理压缩包
    
    Args:
        archive_path: 压缩包路径
        client_name: 客户名称
        output_dir: 输出目录
        watermark_config: 水印配置
        logger: 日志记录器
        
    Returns:
        str: 输出压缩包路径 或 None（失败时）
    """
    try:
        # 创建临时目录
        temp_dir = tempfile.mkdtemp(prefix="watermark_")
        extract_dir = os.path.join(temp_dir, "extracted")
        processed_dir = os.path.join(temp_dir, "processed")
        os.makedirs(extract_dir)
        os.makedirs(processed_dir)
        
        # 解压
        with zipfile.ZipFile(archive_path, 'r') as zf:
            zf.extractall(extract_dir)
        
        logger.info(f"压缩包解压完成: {archive_path}")
        
        # 遍历处理所有文件
        processed_files = []
        for root, dirs, files in os.walk(extract_dir):
            for file in files:
                file_path = os.path.join(root, file)
                file_type = get_file_type(file_path)
                
                if file_type == 'excel':
                    print_progress('processing', filename=file)
                    result = process_excel_file(
                        file_path, client_name, processed_dir, 
                        watermark_config, logger
                    )
                    if result:
                        processed_files.append(result)
                        print_progress('done', filename=file)
                    else:
                        print_progress('error', filename=file)
                        
                elif file_type == 'csv':
                    print_progress('processing', filename=file)
                    result = process_csv_file(
                        file_path, client_name, processed_dir,
                        watermark_config, logger
                    )
                    if result:
                        processed_files.append(result)
                        print_progress('done', filename=file)
                    else:
                        print_progress('error', filename=file)
                else:
                    # 非支持文件，直接复制
                    rel_path = os.path.relpath(file_path, extract_dir)
                    target_path = os.path.join(processed_dir, rel_path)
                    os.makedirs(os.path.dirname(target_path), exist_ok=True)
                    shutil.copy2(file_path, target_path)
        
        # 生成输出压缩包名
        dir_name = os.path.dirname(archive_path)
        file_name = os.path.basename(archive_path)
        name = os.path.splitext(file_name)[0]
        target_dir = output_dir if output_dir else dir_name
        output_archive = os.path.join(target_dir, f"{name}_for_{client_name}.zip")
        
        os.makedirs(os.path.dirname(output_archive) or '.', exist_ok=True)
        
        # 重新打包
        with zipfile.ZipFile(output_archive, 'w', zipfile.ZIP_DEFLATED) as zf:
            for root, dirs, files in os.walk(processed_dir):
                for file in files:
                    file_path = os.path.join(root, file)
                    arc_name = os.path.relpath(file_path, processed_dir)
                    zf.write(file_path, arc_name)
        
        # 清理临时目录
        shutil.rmtree(temp_dir)
        
        logger.info(f"压缩包处理完成: {archive_path} -> {output_archive}")
        return output_archive
        
    except Exception as e:
        logger.error(f"处理压缩包失败: {archive_path}, 错误: {str(e)}")
        return None


def process_single_file(file_path, client_name, output_dir, watermark_config, logger):
    """
    处理单个文件（自动判断类型）
    
    Args:
        file_path: 文件路径
        client_name: 客户名称
        output_dir: 输出目录
        watermark_config: 水印配置
        logger: 日志记录器
        
    Returns:
        str: 输出文件路径 或 None（失败时）
    """
    file_type = get_file_type(file_path)
    
    if file_type == 'excel':
        return process_excel_file(file_path, client_name, output_dir, watermark_config, logger)
    elif file_type == 'csv':
        return process_csv_file(file_path, client_name, output_dir, watermark_config, logger)
    elif file_type == 'archive':
        return process_archive(file_path, client_name, output_dir, watermark_config, logger)
    else:
        logger.error(f"不支持的文件类型: {file_path}")
        return None


def process_batch(file_paths, client_name, output_dir, watermark_config, logger):
    """
    批量处理文件
    
    Args:
        file_paths: 文件路径列表
        client_name: 客户名称
        output_dir: 输出目录
        watermark_config: 水印配置
        logger: 日志记录器
        
    Returns:
        list: 成功处理的输出文件列表
    """
    results = []
    total = len(file_paths)
    
    for idx, file_path in enumerate(file_paths, 1):
        print_progress('processing', current=idx, total=total, 
                      filename=os.path.basename(file_path))
        
        result = process_single_file(
            file_path, client_name, output_dir, 
            watermark_config, logger
        )
        
        if result:
            results.append(result)
            print_progress('done', current=idx, total=total,
                          filename=os.path.basename(file_path))
        else:
            print_progress('error', current=idx, total=total,
                          filename=os.path.basename(file_path))
    
    return results
